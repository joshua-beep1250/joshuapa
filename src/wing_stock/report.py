from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KST = timezone(timedelta(hours=9))


class ReportError(ValueError):
    pass


ALIASES = {
    "sku": ["옵션ID", "상품옵션ID", "SKU ID", "SKU", "판매자상품코드", "판매자SKU"],
    "name": ["상품명", "등록상품명", "노출상품명"],
    "option": ["옵션명", "옵션", "상품옵션명"],
    "stock": ["판매가능재고", "판매가능재고수량", "판매가능수량", "현재재고수량", "현재재고", "재고수량"],
    "sales7": ["최근7일판매수량", "최근7일판매량", "최근7일판매", "7일판매수량", "7일판매량"],
}


def normalize(value) -> str:
    return re.sub(r"[\s_()（）\[\]]", "", str(value or "")).casefold()


@dataclass(frozen=True)
class Product:
    sku: str
    name: str
    option: str
    stock: int
    sales7: int

    @property
    def required(self):
        return self.sales7 * 6

    @property
    def replenish(self):
        return max(self.required - self.stock, 0)


def quantity(value, field, row):
    if value is None or str(value).strip() == "":
        raise ReportError(f"{row}행 {field}: 수량이 비어 있습니다. 0과 누락은 구분해야 합니다.")
    try:
        number = Decimal(str(value).replace(",", "").strip())
    except InvalidOperation:
        raise ReportError(f"{row}행 {field}: 숫자가 아닙니다.") from None
    if not number.is_finite() or number < 0 or number != number.to_integral_value():
        raise ReportError(f"{row}행 {field}: 0 이상의 정수가 필요합니다.")
    return int(number)


def _rows(path: Path, sheet: str | None):
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "cp949"):
            try:
                decoded = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ReportError("CSV 인코딩은 UTF-8 또는 CP949여야 합니다.")
        return list(csv.reader(io.StringIO(decoded)))
    if path.suffix.lower() != ".xlsx":
        raise ReportError(".xlsx 또는 .csv 보고서만 지원합니다. .xls는 .xlsx로 저장해 주세요.")
    with path.open("rb") as source:
        wb = load_workbook(source, read_only=True, data_only=True)
        try:
            if sheet and sheet not in wb.sheetnames:
                raise ReportError(f"시트를 찾을 수 없습니다: {sheet}")
            if not sheet and len(wb.sheetnames) != 1:
                raise ReportError("여러 시트가 있습니다. --sheet로 실제 데이터 시트를 지정하세요.")
            return list(wb[sheet or wb.sheetnames[0]].values)
        finally:
            wb.close()


def parse_report(path: Path, columns: dict | None = None, sheet: str | None = None):
    rows = _rows(path, sheet)
    columns = columns or {}
    if not isinstance(columns, dict) or any(not isinstance(v, str) or not v.strip() for v in columns.values()):
        raise ReportError("columns는 필드 이름과 비어 있지 않은 실제 열 제목의 객체여야 합니다.")
    unknown = set(columns) - set(ALIASES)
    if unknown:
        raise ReportError(f"알 수 없는 열 매핑: {', '.join(sorted(unknown))}")
    header = None
    indexes = None
    for position, row in enumerate(rows[:50]):
        normalized = [normalize(cell) for cell in row]
        candidate = {}
        for field, aliases in ALIASES.items():
            choices = [columns[field]] if field in columns else aliases
            # Prefer a specific alias (e.g. 판매가능재고) over a broad stock column.
            for alias in choices:
                matches = [i for i, cell in enumerate(normalized) if cell == normalize(alias)]
                if len(matches) > 1:
                    raise ReportError(f"동일한 열 제목이 중복됩니다: {alias}")
                if matches:
                    candidate[field] = matches[0]
                    break
        if {"sku", "name", "stock", "sales7"} <= candidate.keys():
            if len(set(candidate.values())) != len(candidate):
                raise ReportError("여러 필드에 동일한 원본 열을 지정할 수 없습니다.")
            header, indexes = position, candidate
            break
    if indexes is None:
        raise ReportError(
            "상품/옵션 ID, 상품명, 판매가능 재고, 최근 7일 판매수량 열을 찾지 못했습니다. "
            "config.local.json의 columns에 실제 열 제목을 지정하세요. "
            "전체 판매량·주문건수·최근 30일 판매량을 7일 판매수량으로 대체하지 마세요."
        )
    products = []
    seen = set()
    for line, row in enumerate(rows[header + 1 :], start=header + 2):
        if not any(value is not None and str(value).strip() for value in row):
            continue

        def cell(field):
            index = indexes.get(field)
            return row[index] if index is not None and index < len(row) else None

        sku = str(cell("sku") if cell("sku") is not None else "").strip()
        if not sku:
            raise ReportError(f"{line}행: 상품/옵션 ID가 없습니다. 합계 행이 있다면 제거해 주세요.")
        if sku in seen:
            raise ReportError(f"{line}행: ID {sku}가 중복됩니다. 창고별 행이나 중복 다운로드를 확인하세요.")
        seen.add(sku)
        name = str(cell("name") or "").strip()
        if not name:
            raise ReportError(f"{line}행: 상품명이 없습니다.")
        products.append(Product(sku, name, str(cell("option") or ""),
                                quantity(cell("stock"), "재고", line),
                                quantity(cell("sales7"), "최근 7일 판매량", line)))
    if not products:
        raise ReportError("상품 데이터가 없습니다. 로그인·필터·다운로드 파일을 확인하세요.")
    headers = {field: str(rows[header][index]) for field, index in indexes.items()}
    return products, headers


def _safe_text(cell, value):
    cell.value = str(value)
    cell.data_type = "s"  # Product names/IDs must never become Excel formulas.


def write_report(products, output: Path, source: Path, headers: dict, scope: str):
    if output.resolve() == source.resolve():
        raise ReportError("원본 보고서에 덮어쓸 수 없습니다.")
    if output.exists():
        raise ReportError(f"출력 파일이 이미 있습니다: {output}. 다른 이름을 지정하세요.")
    wb = Workbook()
    ws = wb.active
    ws.title = "재고보충계획"
    titles = ["상품·옵션 ID", "상품명", "옵션명", "판매가능 재고", "최근 7일 판매수량",
              "일평균 판매수량", "6주 필요재고", "추가 확보수량", "재고 소진 예상일수", "상태"]
    ws.append(titles)
    for p in sorted(products, key=lambda p: (-p.replenish, p.sku)):
        status = "판매없음" if not p.sales7 else "품절" if not p.stock else "확보필요" if p.replenish else "충분"
        ws.append([None, None, None, p.stock, p.sales7, p.sales7 / 7,
                   p.required, p.replenish, p.stock * 7 / p.sales7 if p.sales7 else None, status])
        row = ws.max_row
        for column, value in enumerate((p.sku, p.name, p.option), 1):
            _safe_text(ws.cell(row, column), value)
        for column in (6, 9):
            ws.cell(row, column).number_format = "0.00"
        if p.replenish:
            color = "FCE4D6" if p.stock else "FFC7CE"
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=color)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    for i, width in enumerate((24, 45, 32, 19, 23, 22, 19, 19, 24, 16), 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    summary = wb.create_sheet("요약·계산기준")
    now = datetime.now(KST).isoformat(timespec="seconds")
    details = [
        ("보고서 생성시각 (한국)", now), ("원본 파일", source.name),
        ("조회 범위·기간 확인", scope), ("최근 7일의 정의", "WING 원본 지표 기준. 원본의 집계기간·갱신시각을 확인하세요."),
        ("상품·옵션 수", len(products)), ("현재 재고 합계", sum(p.stock for p in products)),
        ("7일 판매수량 합계", sum(p.sales7 for p in products)),
        ("추가 확보수량 합계", sum(p.replenish for p in products)),
        ("확보 필요 상품·옵션 수", sum(p.replenish > 0 for p in products)),
        ("6주 필요재고 계산", "최근 7일 판매수량 × 6 (42일)"),
        ("추가 확보수량 계산", "max(6주 필요재고 - 판매가능 재고, 0)"),
        ("가정", "현재 판매 속도 유지. 입고 예정·안전재고·리드타임·계절성 미반영."),
        ("품절 유의", "품절 기간의 판매량은 실제 수요보다 작을 수 있습니다."),
        ("원본 기록", "원본 다운로드와 함께 보관. 현재 재고는 다운로드 시점 값."),
    ] + [(f"원본 열: {key}", value) for key, value in headers.items()]
    for key, value in details:
        summary.append([key, value])
        if isinstance(value, str):
            _safe_text(summary.cell(summary.max_row, 2), value)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 95
    for sheet in wb:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation prevents accidental replacement of another report.
    with output.open("xb") as destination:
        wb.save(destination)
    return output
