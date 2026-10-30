from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from wing_stock.cli import main
from wing_stock.report import ReportError, parse_report, write_report


HEADER = "옵션ID,상품명,옵션명,판매가능 재고,최근 7일 판매수량\n"


def csv_file(tmp_path, body, header=HEADER):
    path = tmp_path / "input.csv"
    path.write_text(header + body, encoding="utf-8-sig")
    return path


def test_replenishment_and_excel(tmp_path):
    path = csv_file(tmp_path, "001,수건,흰색,20,10\n002,컵,빨강,0,3\n003,함,대형,100,2\n004,양말,검정,12,0\n")
    products, headers = parse_report(path)
    assert [p.required for p in products] == [60, 18, 12, 0]
    assert [p.replenish for p in products] == [40, 18, 0, 0]
    out = tmp_path / "result.xlsx"
    write_report(products, out, path, headers, "전체 상품, WING 최근 7일")
    wb = load_workbook(out, data_only=True)
    ws = wb["재고보충계획"]
    assert ws["A2"].value == "001"  # Leading zeros retained.
    assert ws["H2"].value == 40
    assert ws["H3"].value == 18
    assert ws["J3"].value == "품절"
    assert ws["I5"].value is None  # No division by zero.
    assert ws.auto_filter.ref == "A1:J5"
    assert dict(wb["요약·계산기준"].values)["추가 확보수량 합계"] == 58
    wb.close()


@pytest.mark.parametrize("bad", ["", "-1", "1.5", "NaN", "Infinity", "판매 3건", "-"])
def test_bad_or_missing_quantity_fails(tmp_path, bad):
    path = csv_file(tmp_path, f"a,상품,옵션,10,{bad}\n")
    with pytest.raises(ReportError):
        parse_report(path)


def test_duplicate_option_fails(tmp_path):
    path = csv_file(tmp_path, "1,상품,옵션,10,2\n1,상품,옵션,20,2\n")
    with pytest.raises(ReportError, match="중복"):
        parse_report(path)


def test_missing_sales7_never_defaults_to_zero(tmp_path):
    path = csv_file(tmp_path, "1,상품,옵션,10,12\n", HEADER.replace("최근 7일 판매수량", "최근 30일 판매수량"))
    with pytest.raises(ReportError, match="찾지 못"):
        parse_report(path)


def test_empty_report_fails(tmp_path):
    with pytest.raises(ReportError, match="없습니다"):
        parse_report(csv_file(tmp_path, ""))


def test_xlsx_preamble_sheet_mapping_and_commas(tmp_path):
    wb = Workbook()
    wb.active.title = "안내"
    ws = wb.create_sheet("데이터")
    ws.append(["보고서 안내"])
    ws.append(["식별자", "품명", "가용수량", "최근7일실판매"])
    ws.append(["000123", "품목", "1,000", "200"])
    path = tmp_path / "source.xlsx"
    wb.save(path)
    with pytest.raises(ReportError, match="여러 시트"):
        parse_report(path)
    products, _ = parse_report(path, {"sku": "식별자", "name": "품명", "stock": "가용수량", "sales7": "최근7일실판매"}, "데이터")
    assert products[0].replenish == 200
    assert products[0].sku == "000123"


def test_formula_like_text_is_literal_and_original_is_unchanged(tmp_path):
    source = csv_file(tmp_path, "=1+1,=HYPERLINK(1),+옵션,2,1\n")
    original = source.read_bytes()
    products, headers = parse_report(source)
    out = tmp_path / "safe.xlsx"
    write_report(products, out, source, headers, "테스트")
    wb = load_workbook(out)
    assert wb.active["A2"].data_type == "s"
    assert wb.active["B2"].data_type == "s"
    wb.close()
    assert source.read_bytes() == original
    with pytest.raises(ReportError, match="이미"):
        write_report(products, out, source, headers, "테스트")
    with pytest.raises(ReportError, match="원본"):
        write_report(products, source, source, headers, "테스트")


def test_cli_real_file_and_error_exit(tmp_path, capsys):
    source = Path(__file__).parents[1] / "examples" / "inventory.csv"
    output = tmp_path / "cli.xlsx"
    assert main(["--input", str(source), "--scope", "샘플 최근 7일", "--output", str(output)]) == 0
    assert output.exists()
    assert "58개" in capsys.readouterr().out
    assert main(["--input", str(source)]) == 1
    assert "--scope" in capsys.readouterr().err


def test_cp949(tmp_path):
    path = tmp_path / "cp949.csv"
    path.write_bytes((HEADER + "1,상품,옵션,1,1\n").encode("cp949"))
    assert parse_report(path)[0][0].name == "상품"


def test_specific_available_stock_over_total_stock(tmp_path):
    path = csv_file(tmp_path, "1,상품,20,7,2\n", "옵션ID,상품명,재고수량,판매가능수량,최근7일판매량\n")
    assert parse_report(path)[0][0].stock == 7


def test_same_column_cannot_be_stock_and_sales(tmp_path):
    path = csv_file(tmp_path, "1,상품,옵션,10,2\n")
    with pytest.raises(ReportError, match="동일한 원본 열"):
        parse_report(path, {"sales7": "판매가능 재고"})


def test_default_browser_manual_download_mode(tmp_path, monkeypatch):
    opened = []
    monkeypatch.setattr("webbrowser.open", lambda url, new: opened.append(url) or True)
    source = Path(__file__).parents[1] / "examples" / "inventory.csv"
    answers = iter([f'"{source}"', "전체 상품, 최근 7일 (테스트)"])
    monkeypatch.setattr("builtins.input", lambda _: next(answers))
    output = tmp_path / "manual.xlsx"
    assert main(["--manual-download", "--output", str(output)]) == 0
    assert output.is_file()
    assert opened == ["https://wing.coupang.com/"]
